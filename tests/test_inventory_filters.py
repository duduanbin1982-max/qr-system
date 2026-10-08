import os
import sqlite3
import json
import uuid
import csv
from io import BytesIO, StringIO

from openpyxl import load_workbook

import pytest

from factories import TEST_HASH, TEST_PASS
from modules.db import get_db
from tests.inventory_product_helpers import seed_product_inventory_scenario


@pytest.fixture
def db():
    connection = sqlite3.connect(os.environ["DB_PATH"])
    connection.row_factory = sqlite3.Row
    try:
        yield connection
    finally:
        connection.close()


def _permission_headers(client, permissions):
    suffix = uuid.uuid4().hex[:8]
    username = f"inventory_filter_{suffix}"
    role_code = f"inventory_filter_{suffix}"
    with client.application.app_context():
        connection = get_db()
        role_id = connection.execute(
            "INSERT INTO roles (name,code,description,permissions,status,level) "
            "VALUES (?,?,?,?,'active',1)",
            ("Inventory Filter", role_code, "pytest inventory role", json.dumps(permissions)),
        ).lastrowid
        user_id = connection.execute(
            "INSERT INTO users "
            "(username,password,name,role,status,password_version,employee_no) "
            "VALUES (?,?,?,?,'active',2,?)",
            (username, TEST_HASH, "Inventory Filter", role_code, f"INV-{suffix}"),
        ).lastrowid
        connection.execute(
            "INSERT INTO user_roles(user_id,role_id) VALUES (?,?)",
            (user_id, role_id),
        )
        connection.commit()
    response = client.post(
        "/api/auth/login", json={"username": username, "password": TEST_PASS}
    )
    assert response.status_code == 200
    return {"Authorization": f"Bearer {response.get_json()['user']['token']}"}


def test_order_inventory_filters_support_specification_quality_and_pagination(
    client, auth_headers, db
):
    scenario = seed_product_inventory_scenario(db, specs=("标准", "加厚"))

    options = client.get("/api/inventory/filter-options", headers=auth_headers)
    assert options.status_code == 200
    payload = options.get_json()
    assert {item["value"] for item in payload["specifications"]} >= {"标准", "加厚"}
    assert payload["quality_statuses"]

    filtered = client.get(
        "/api/inventory?specification=加厚&quality_status=qualified&page=1&limit=1",
        headers=auth_headers,
    )
    assert filtered.status_code == 200
    body = filtered.get_json()
    assert body["total"] == 1
    assert body["limit"] == 1
    assert len(body["items"]) == 1
    assert body["items"][0]["specification"] == "加厚"
    assert body["items"][0]["id"] == scenario["inventory_ids"][1]

    stats = client.get(
        "/api/inventory/stats?specification=加厚&quality_status=qualified",
        headers=auth_headers,
    )
    assert stats.status_code == 200
    summary = stats.get_json()
    assert summary["total_items"] == 1
    assert summary["total_quantity"] == 7
    db.execute(
        "UPDATE products SET price=12.5 WHERE id=?", (scenario["product_id"],)
    )
    db.commit()
    valued = client.get(
        "/api/inventory/stats?specification=加厚&quality_status=qualified",
        headers=auth_headers,
    ).get_json()
    assert valued["total_value"] == 87.5


def test_order_inventory_empty_specification_filter(client, auth_headers, db):
    db.execute(
        "INSERT INTO inventory(product_model,product_name,specification,quantity) "
        "VALUES('EMPTY-SPEC-1','空规格','',3)"
    )
    db.commit()
    response = client.get(
        "/api/inventory?specification=__empty__", headers=auth_headers
    )
    assert response.status_code == 200
    body = response.get_json()
    assert body["total"] == 1
    assert body["items"][0]["specification"] == ""


def test_default_inventory_stats_subtract_frozen_quantity(
    client, auth_headers, db
):
    db.execute(
        "INSERT INTO inventory(product_model,product_name,quantity,reserved,"
        "frozen_quantity,safe_stock) VALUES('FROZEN-STATS-1','冻结测试',10,0,4,7)"
    )
    db.commit()

    items = client.get("/api/inventory", headers=auth_headers).get_json()
    stats = client.get("/api/inventory/stats", headers=auth_headers).get_json()

    assert items["total"] == 1
    assert items["items"][0]["available_quantity"] == 6
    assert items["items"][0]["is_low"] == 1
    assert stats["total_items"] == 1
    assert stats["low_stock"] == 1


def test_inventory_export_uses_the_same_filters_as_the_order_view(
    client, auth_headers, db
):
    seed_product_inventory_scenario(db, specs=("标准", "加厚"))

    response = client.get(
        "/api/inventory/export?specification=加厚&quality_status=qualified",
        headers=auth_headers,
    )

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.data), read_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    assert len(rows) == 2
    assert rows[1][4] == "加厚"


def test_inventory_csv_export_uses_filters_and_preserves_quoted_values(client, auth_headers, db):
    db.execute(
        "INSERT INTO inventory(product_model,product_name,quantity,location,"
        "specification,quality_status,unit,remark) VALUES(?,?,?,?,?,?,?,?)",
        ("00000001", '产品,一', 4, "A-01", "加厚", "qualified", "件", '说明"一'),
    )
    db.commit()
    response = client.get(
        "/api/inventory/export.csv?specification=加厚&quality_status=qualified",
        headers=auth_headers,
    )
    assert response.status_code == 200
    assert response.content_type.startswith("text/csv")
    assert response.data.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(StringIO(response.data.decode("utf-8-sig"))))
    assert rows[0][0:5] == ["产品名称", "订单号", "客户", "产品型号", "规格"]
    assert rows[1][0] == "产品,一"
    assert rows[1][3] == "00000001"


def test_inventory_export_is_not_truncated_at_the_page_limit(
    client, auth_headers, db
):
    db.executemany(
        "INSERT INTO inventory(product_model,product_name,quantity,unit_cost) "
        "VALUES(?,?,1,1)",
        [(f"BULK-{index:04d}", "批量导出") for index in range(501)],
    )
    db.commit()

    response = client.get("/api/inventory/export", headers=auth_headers)

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.data), read_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    assert len(rows) == 502


def test_order_inventory_sort_is_stable_and_inherited_by_export(client, auth_headers, db):
    db.executemany(
        "INSERT INTO inventory(product_model,product_name,quantity,updated_at) VALUES(?,?,?,?)",
        [
            ("SORT-LOW", "排序低", 2, "2026-01-01 00:00:00"),
            ("SORT-HIGH", "排序高", 9, "2026-01-01 00:00:00"),
        ],
    )
    db.commit()

    response = client.get(
        "/api/inventory?sort_by=quantity&sort_dir=desc&limit=10",
        headers=auth_headers,
    )
    assert response.status_code == 200
    items = response.get_json()["items"]
    assert [item["product_model"] for item in items[:2]] == ["SORT-HIGH", "SORT-LOW"]

    exported = client.get(
        "/api/inventory/export?sort_by=quantity&sort_dir=desc",
        headers=auth_headers,
    )
    assert exported.status_code == 200
    rows = list(load_workbook(BytesIO(exported.data), read_only=True).active.iter_rows(values_only=True))
    assert rows[1][3] == "SORT-HIGH"


def test_inventory_exports_require_the_export_permission(client):
    view_only = _permission_headers(client, ["inventory:view"])
    assert client.get("/api/inventory", headers=view_only).status_code == 200
    assert client.get("/api/inventory/export", headers=view_only).status_code == 403
    assert client.get("/api/inventory/export.csv", headers=view_only).status_code == 403
    assert (
        client.get("/api/inventory/product-groups/export", headers=view_only).status_code
        == 403
    )
    assert client.get("/api/inventory/product-groups/export.csv", headers=view_only).status_code == 403
